"""Automatic ingestion for materials uploaded from the FOF workbench.

The ingestion result is deliberately reviewable: product entities and NAV
observations created here remain pending until a user confirms or reviews them.

Parsing strategy per file type:
- Image (weekly report): table OCR for disclosed metrics, then chart_extractor
  pixel tracing (VLM when configured, else auto-color CV) for NAV curves.
- PDF: text-layer extraction for context + chart pipeline (render → detect →
  trace) for NAV curves.
- XLSX/CSV: long format (date + nav column) or wide format (date + one
  column per product).
"""

from __future__ import annotations

import asyncio
import calendar
import json
import logging
import os
import re
import zipfile
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from sqlalchemy import select

from app.database import SessionFactory
from app.models import DocumentFragment, NavCandidateVersion, NavObservation, ProductEntity
from app.services import product_store as store
from app.services.multi_product_report import extract_multi_product_report

logger = logging.getLogger(__name__)

_IMAGE_TYPES = {"image/jpeg", "image/png", "image/webp"}
_TABLE_SUFFIXES = {".csv", ".xlsx", ".xls"}
_OFFICE_SUFFIXES = {".docx", ".pptx"}


def _normalise_observation_date(value: Any) -> str | None:
    """Return an ISO date only when the source contains a complete date.

    Visual/OCR output sometimes contains axis labels such as ``21/11`` or
    ``10月``.  Those values do not contain a year, so assigning one would turn
    a chart-reading artefact into fabricated NAV history.  Keep the source
    evidence, but exclude that individual point from the NAV series instead
    of failing the complete material-ingestion job.
    """
    if isinstance(value, date):
        parsed = value
    else:
        if not isinstance(value, str):
            return None
        token = value.strip().replace("/", "-").replace(".", "-")
        if len(token) == 8 and token.isdigit():
            token = f"{token[:4]}-{token[4:6]}-{token[6:]}"
        try:
            parsed = date.fromisoformat(token)
        except ValueError:
            return None
    # A chart OCR artefact can be syntactically valid while still being an
    # impossible NAV date (e.g. 2078-12-02).  NAV history must not contain
    # future observations merely because an axis label was misread.
    if parsed < date(1990, 1, 1) or parsed > date.today():
        return None
    return parsed.isoformat()


def _valid_nav_points(points: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Keep only finite positive NAV values with complete plausible dates."""
    result: list[dict[str, Any]] = []
    for point in _normalise_nav_points(points):
        try:
            value = float(point.get("nav"))
        except (TypeError, ValueError):
            continue
        if not np.isfinite(value) or value <= 0 or value > 1_000_000:
            continue
        result.append({**point, "nav": value})
    return result


def _normalise_nav_points(points: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Discard malformed visual dates without discarding the uploaded file."""
    normalised: list[dict[str, Any]] = []
    for point in points:
        observation_date = _normalise_observation_date(point.get("observation_date"))
        nav = point.get("nav")
        if observation_date is None or nav is None:
            continue
        normalised.append({**point, "observation_date": observation_date})
    return normalised


def _visual_method() -> str:
    """Return the actual image method used for provenance labels."""
    provider_type = os.getenv("VLM_PROVIDER", "dashscope")
    configured = bool(os.getenv("DASHSCOPE_API_KEY", "")) if provider_type == "dashscope" else bool(os.getenv("VLM_BASE_URL", ""))
    return "VLM 结构识别 + CV 像素追踪" if configured else "CV 像素追踪（无 VLM）"


def ingest_uploaded_material(
    file_id: str,
    content: bytes,
    filename: str,
    mime_type: str | None,
    material_nature: str | None = None,
    product_name_hint: str | None = None,
    manager_name_hint: str | None = None,
) -> None:
    """Parse one uploaded file and persist reviewable evidence in the KB."""
    with SessionFactory() as session:
        store.update_file_status(session, file_id, "processing")
        try:
            suffix = Path(filename).suffix.lower()
            context_only = material_nature in {"CTA策略介绍", "CTA研究/市场资料"}
            visual_nav_allowed = not context_only
            if not visual_nav_allowed:
                audit = _ingest_context_material(session, file_id, content, filename, suffix, mime_type)
            elif mime_type in _IMAGE_TYPES or suffix in {".jpg", ".jpeg", ".png", ".webp"}:
                audit = _ingest_report_image(
                    session, file_id, content, filename,
                    product_name_hint=product_name_hint,
                    manager_name_hint=manager_name_hint,
                )
            elif suffix in _TABLE_SUFFIXES:
                rows = _ingest_nav_table(session, file_id, content, filename, suffix)
                audit = {"configured": False, "attempted": False, "succeeded": False, "methods": ["Pandas 表格解析"], "nav_rows": rows}
            elif suffix in _OFFICE_SUFFIXES:
                rows = _ingest_office_document(session, file_id, content, filename, suffix)
                audit = {"configured": False, "attempted": False, "succeeded": False, "methods": ["DOCX/PPTX XML 提取"], "nav_rows": rows}
            elif mime_type == "application/pdf" or suffix == ".pdf":
                audit = _ingest_pdf(
                    session,
                    file_id,
                    content,
                    filename,
                    product_name_hint=product_name_hint,
                )
            else:
                raise ValueError("暂不支持该文件类型的自动解析；请上传 PDF、图片、XLSX、CSV、DOCX 或 PPTX")
        except Exception as error:  # Background jobs must surface a concise, user-visible status.
            # Preserve the failure, but also persist the configured extraction
            # path so the UI can distinguish a VLM outage from a bad document.
            audit = locals().get("audit") or _empty_extraction_audit()
            audit = {**audit, "status": "failed", "error": str(error)[:300]}
            store.update_file_extraction_audit(session, file_id, audit)
            store.update_file_status(session, file_id, "failed", str(error)[:300])
        else:
            store.update_file_extraction_audit(session, file_id, audit)
            candidate_count = session.execute(
                select(NavCandidateVersion.id).where(NavCandidateVersion.source_file_id == file_id)
            ).scalars().all()
            nav_count = session.execute(
                select(NavObservation.id).where(NavObservation.source_file_id == file_id)
            ).scalars().all()
            if candidate_count or nav_count:
                store.update_file_status(session, file_id, "completed")
            else:
                charts = session.execute(
                    select(DocumentFragment.id).where(
                        DocumentFragment.file_id == file_id,
                        DocumentFragment.fragment_type.in_(("chart", "chart_traced")),
                    )
                ).scalars().all()
                detail = "已定位图表区域，但未提取到可校准的净值曲线；请打开原图选择曲线颜色或手动框选坐标轴。" if charts else "未定位到可提取的净值图表；资料已保留，可在原图中手动框选。"
                store.update_file_status(session, file_id, "completed_no_nav", detail)


def _ingest_context_material(
    session: Any,
    file_id: str,
    content: bytes,
    filename: str,
    suffix: str,
    mime_type: str | None,
) -> dict[str, Any]:
    """Retain strategy/research evidence without pretending charts are fund NAV."""
    if suffix in _OFFICE_SUFFIXES:
        _ingest_office_document(session, file_id, content, filename, suffix)
        method = "DOCX/PPTX XML 提取（上下文资料）"
    elif mime_type == "application/pdf" or suffix == ".pdf":
        import fitz

        document = fitz.open(stream=content, filetype="pdf")
        text_parts: list[str] = []
        text_length = 0
        for page_index in range(min(document.page_count, 80)):
            page_text = document.load_page(page_index).get_text("text")
            if page_text:
                text_parts.append(page_text)
                text_length += len(page_text)
            if text_length >= 20000:
                break
        document.close()
        text = "\n".join(text_parts).strip()
        if not text:
            raise ValueError("上下文 PDF 无可读取文本层；未将图表误作产品净值")
        store.add_fragment(
            session,
            file_id=file_id,
            fragment_type="text",
            content_text=text[:20000],
            content_data={"method": "PDF 文本层提取（上下文资料）", "review_status": "pending"},
            ocr_confidence=0.9,
            product_id=None,
        )
        method = "PDF 文本层提取（上下文资料）"
    else:
        # The classification manifest has already established that this is
        # context, not a product-performance image. Keep the source reviewable
        # and avoid an expensive VLM call that cannot produce formal NAV.
        store.add_fragment(
            session,
            file_id=file_id,
            fragment_type="image",
            content_data={"method": "分类清单分流（上下文图片）", "review_status": "pending"},
            product_id=None,
        )
        method = "分类清单分流（上下文图片）"
    return {"configured": False, "attempted": False, "succeeded": False, "methods": [method]}

def _empty_extraction_audit() -> dict[str, Any]:
    provider = os.getenv("VLM_PROVIDER", "dashscope")
    model = os.getenv("VLM_MODEL", "qwen3-vl-flash")
    configured = bool(os.getenv("DASHSCOPE_API_KEY", "")) if provider == "dashscope" else bool(os.getenv("VLM_BASE_URL", ""))
    return {"configured": configured, "attempted": False, "succeeded": False, "provider": provider, "model": model}


def _merge_extraction_audit(*audits: dict[str, Any]) -> dict[str, Any]:
    """Combine per-page/per-curve provenance into one file-level audit."""
    valid = [item for item in audits if item]
    if not valid:
        return _empty_extraction_audit()
    first = valid[0]
    errors: list[str] = []
    for item in valid:
        if item.get("error"):
            errors.append(str(item["error"]))
        errors.extend(str(error) for error in item.get("errors", []) if error)
    methods: list[str] = []
    for item in valid:
        for method in item.get("methods", []):
            if method not in methods:
                methods.append(method)
    return {
        "configured": any(bool(item.get("configured")) for item in valid),
        "attempted": any(bool(item.get("attempted")) for item in valid),
        "succeeded": any(bool(item.get("succeeded")) for item in valid),
        "provider": next((item.get("provider") for item in valid if item.get("provider")), first.get("provider")),
        "model": next((item.get("model") for item in valid if item.get("model")), first.get("model")),
        "calls": sum(int(item.get("calls", 0)) for item in valid),
        "successful_calls": sum(int(item.get("successful_calls", 0)) for item in valid),
        "failed_calls": sum(int(item.get("failed_calls", 0)) for item in valid),
        "errors": errors[:20],
        "methods": methods,
    }


def _ingest_report_image(
    session: Any,
    file_id: str,
    content: bytes,
    filename: str = "",
    *,
    product_name_hint: str | None = None,
    manager_name_hint: str | None = None,
) -> dict[str, Any]:
    report = extract_multi_product_report(content)
    # Even when OCR cannot read a product table or trace a curve, keep a
    # reviewable candidate in the workbench.  A completed file must never look
    # like it simply disappeared from the user's product queue.
    # The filename is a valuable identity hint (e.g. manager_product.png).
    # Prefer an existing filename candidate for a one-product report so a
    # later OCR title such as "低波CTA" enriches the same product instead of
    # creating a second, disconnected product record.
    multi_product_material = _is_multi_product_material(filename, product_name_hint)
    manifest_name = _normalise_manifest_product_name(product_name_hint)
    filename_name = manifest_name or _product_name_from_filename(filename)
    filename_product = None
    if filename_name and not multi_product_material:
        filename_product = _find_or_create_product(
            session,
            filename_name,
            report.product_identity.strategy,
            None,
            None,
            manager_name_hint or _manager_name_from_filename(filename),
        )
    fallback_product = filename_product
    if not report.disclosed_metrics:
        # Product identity is independent evidence.  A factsheet may have a
        # perfectly readable title while its KPI row is incomplete; do not
        # fall back to an opaque file stem in that case.
        # A filename such as "低波CTA策略" is a material/strategy label, not
        # proof of a distinct fund.  Keep the file reviewable but leave it
        # unbound until the user selects the concrete product.  This prevents
        # one factsheet from producing both a strategy pseudo-product and its
        # actual product.
        identity_name = "" if multi_product_material else (report.product_identity.product_name or "").strip()
        fallback_name = identity_name if _looks_like_product_name(identity_name) else ""
        if fallback_product is None and fallback_name:
            fallback_product = _find_or_create_product(
                session,
                fallback_name,
                report.product_identity.strategy,
                None,
                None,
                report.product_identity.manager_name,
            )
    # Cheap, explicit routing before the costly pixel trace.  A multi-product
    # page must go through curve binding; a page without a product identity is
    # retained as material evidence rather than spending time guessing a line.
    if multi_product_material or len(report.disclosed_metrics) > 1:
        material_lane = "multi_product"
        trace_allowed = False
        triage_reason = "文件标记为多产品资料，需先确认曲线与产品归属"
    elif fallback_product is None:
        material_lane = "non_product_or_unbound"
        trace_allowed = False
        triage_reason = "未获得可绑定的单产品身份，未自动追踪曲线"
    else:
        material_lane = "single_product_nav"
        trace_allowed = True
        triage_reason = "单产品身份可绑定，进入净值曲线追踪"
    store.add_fragment(
        session,
        file_id=file_id,
        fragment_type="image",
        content_data={
            "method": _visual_method(),
            "warnings": report.warnings,
            "product_identity": report.product_identity.model_dump(mode="json"),
            "curve_candidates": [item.model_dump() for item in report.product_curve_candidates],
            "vlm_layout_attempted": report.vlm_layout_attempted,
            "vlm_layout_used": report.vlm_layout_used,
            "confidence": report.product_identity.confidence or None,
            "material_lane": material_lane,
            "trace_allowed": trace_allowed,
            "triage_reason": triage_reason,
            "workflow": {
                "stage": "trace" if trace_allowed else "review_binding",
                "identity_source": "classification_manifest" if manifest_name else report.product_identity.method,
                "next_action": "trace_chart" if trace_allowed else "confirm_product_binding",
                "retryable": True,
            },
        },
        ocr_confidence=report.product_identity.confidence or None,
        product_id=fallback_product.id if fallback_product else None,
    )
    products_by_source_id: dict[str, ProductEntity] = {}
    for metric in report.disclosed_metrics:
        name = metric.product_name or metric.product_id
        # A one-product image whose filename carries a concrete product name
        # has stronger identity evidence than OCR of labels such as “单位” or
        # “累计”.  Do not create a product from that noisy label.
        if multi_product_material:
            # A page-level multi-product hint is not enough to bind its only
            # OCR row to one fund; wait for explicit product/curve evidence.
            product = None if len(report.disclosed_metrics) == 1 else _find_or_create_product(
                session, name, metric.strategy, metric.start_date, metric.end_date,
                report.product_identity.manager_name,
            )
        elif len(report.disclosed_metrics) == 1:
            product = filename_product
            if product is None and _looks_like_product_name(name):
                product = _find_or_create_product(
                    session, name, metric.strategy, metric.start_date, metric.end_date,
                    report.product_identity.manager_name,
                )
        else:
            product = _find_or_create_product(
                session, name, metric.strategy, metric.start_date, metric.end_date,
                report.product_identity.manager_name,
            )
        if product is filename_product and product is not None:
            if metric.strategy and not product.strategy:
                product.strategy = metric.strategy
            if metric.start_date and product.inception_date is None:
                product.inception_date = metric.start_date
            if metric.end_date and (product.close_date is None or metric.end_date > product.close_date):
                product.close_date = metric.end_date
            session.commit()
        fragment = store.add_fragment(
            session,
            file_id=file_id,
            fragment_type="table",
            content_text=f"{name}：区间 {metric.start_date} 至 {metric.end_date}",
            content_data={**metric.model_dump(mode="json"), "method": "OCR/表格指标解析"},
            ocr_confidence=0.65,
            product_id=product.id if product is not None else None,
        )
        disclosed_facts = {
            "disclosed_cumulative_return": metric.cumulative_return,
            "disclosed_annualized_return": metric.annualized_return,
        }
        if metric.maximum_drawdown_disclosed:
            disclosed_facts["disclosed_maximum_drawdown"] = metric.maximum_drawdown
        if product is not None:
            products_by_source_id[metric.product_id] = product
            for field_name, value in disclosed_facts.items():
                store.add_fact(
                    session,
                    product_id=product.id if product else None,
                    field_name=field_name,
                    field_value=str(value),
                    source_file_id=file_id,
                    source_fragment_id=fragment.id,
                    confidence=0.65,
                )

    ordered_products = [products_by_source_id.get(m.product_id) for m in report.disclosed_metrics]
    if not ordered_products and fallback_product is not None:
        ordered_products = [fallback_product]
    for curve in report.product_curve_candidates:
        product = products_by_source_id.get(curve.product_id or "")
        curve_is_bound = product is not None or (
            len(report.disclosed_metrics) == 1 and not multi_product_material
        )
        store.add_fragment(
            session,
            file_id=file_id,
            fragment_type="chart",
            bbox={"left_ratio": curve.left_ratio, "top_ratio": curve.top_ratio, "right_ratio": curve.right_ratio, "bottom_ratio": curve.bottom_ratio},
            content_data={
                "curve_index": curve.curve_index,
                "source_product_id": curve.product_id,
                "method": _visual_method(),
                "binding_status": "matched" if curve_is_bound else "unmatched",
                "binding_reason": None if curve_is_bound else "多产品曲线未返回可验证的产品 ID",
                "binding_confidence": curve.binding_confidence,
                "binding_evidence": curve.binding_evidence,
                "legend_label": curve.legend_label,
                "layout_product_name": curve.layout_product_name,
                "color_hex": curve.color_hex,
            },
            ocr_confidence=0.45,
            product_id=product.id if product else None,
        )

    # --- Curve tracing: convert pixel curves into NAV observations ---
    trace_audits: list[dict[str, Any]] = []
    if trace_allowed and report.product_curve_candidates:
        trace_audits = _trace_image_curves(
            session, file_id, content, report.product_curve_candidates,
            ordered_products, filename, report.disclosed_metrics,
        )
    elif not report.disclosed_metrics and fallback_product is not None and trace_allowed:
        # No table, no candidates — treat the whole image as one NAV chart.
        trace_audits = [_trace_whole_image(session, file_id, content, fallback_product.standard_name if fallback_product else (Path(filename).stem if filename else "未命名产品"))]
    # A few red/green table cells can be mistaken for a short chart trace.
    # Such a low-confidence candidate must not block the stronger monthly
    # return-table reconstruction lane.
    has_reliable_candidate = session.execute(
        select(NavCandidateVersion.id).where(
            NavCandidateVersion.source_file_id == file_id,
            NavCandidateVersion.confidence >= 0.65,
        )
    ).scalar() is not None
    if trace_allowed and not has_reliable_candidate and fallback_product is not None:
        recovery_points, recovery_audit = _recover_disclosed_nav_table(content)
        trace_audits.append(recovery_audit)
        recovery_method = "VLM 披露净值表读取"
        recovery_confidence = 0.9
        if len(recovery_points) < 3:
            recovery_points, recovery_audit = _recover_monthly_nav_from_return_table(content)
            trace_audits.append(recovery_audit)
            recovery_method = "VLM 月收益表读取 + 确定性复利重建"
            recovery_confidence = 0.72
        if len(recovery_points) >= 3:
            fragment = store.add_fragment(
                session,
                file_id=file_id,
                fragment_type="table",
                content_text=f"披露表格恢复 {len(recovery_points)} 个候选净值点",
                content_data={
                    "method": recovery_method,
                    "review_status": "pending",
                    "workflow_stage": "table_recovery",
                    "num_points": len(recovery_points),
                },
                ocr_confidence=recovery_confidence,
                product_id=fallback_product.id,
            )
            candidate = store.create_nav_candidate_version(
                session,
                fallback_product.id,
                file_id,
                recovery_points,
                source_fragment_id=fragment.id,
                frequency=_infer_frequency([
                    date.fromisoformat(point["observation_date"]) for point in recovery_points
                ]),
                confidence=recovery_confidence,
            )
            fragment.content_data = {**(fragment.content_data or {}), "candidate_version_id": candidate.id}
            session.commit()
    audit = _merge_extraction_audit(report.vlm_audit, *trace_audits)
    audit["workflow"] = {
        "identity": "resolved" if fallback_product else "needs_review",
        "layout": "resolved" if report.product_curve_candidates else "not_found",
        "binding": "resolved" if trace_allowed else "needs_review",
        "trace": "candidate_created" if any(
            session.execute(select(NavCandidateVersion.id).where(NavCandidateVersion.source_file_id == file_id)).scalars().all()
        ) else "needs_calibration",
        "next_action": "review_candidate" if session.execute(
            select(NavCandidateVersion.id).where(NavCandidateVersion.source_file_id == file_id)
        ).scalar() else ("confirm_product_binding" if not trace_allowed else "recover_chart_calibration"),
    }
    return audit


_MONTHLY_RETURN_TABLE_PROMPT = """这是一份私募产品业绩报告。请读取明确印刷的逐月收益率表，并读取每行最右侧的全年/今年以来合计以及页面印刷的当前净值，用于交叉校验。
严格输出 JSON：{"monthly_returns":[{"year":2024,"month":1,"return_pct":1.2}],"annual_returns":[{"year":2024,"return_pct":12.5}],"latest_nav":1.3922}
规则：
1. 每项必须来自表格中“年份行 × 月份列”的交叉单元格；空白单元格不要输出。
2. return_pct 使用百分数数值，例如 -1.3% 输出 -1.3。
3. 表格第一列若为“周收益”，它不是月份，绝对不能作为1月或其他月份；月份只对应1月到12月的明确列头。
4. annual_returns 只读取该年份行最右侧“全年”列；latest_nav 只读取“当前净值/最新净值”旁的印刷数字。
5. 无可靠逐月收益表时输出 {"monthly_returns":[],"annual_returns":[],"latest_nav":null}；只输出 JSON。"""


_DISCLOSED_NAV_TABLE_PROMPT = """这是一份私募产品净值/业绩报告。请只读取页面明确印刷的历史净值明细表，不要从曲线估算数据。
严格输出 JSON：{"nav_rows":[{"date":"2026-07-03","unit_nav":1.9810}],"latest_unit_nav":1.9810}
规则：
1. nav_rows 只允许来自同时印刷了完整年月日和“单位净值”的表格行；不要使用复权净值、累计净值或净值变动列。
2. 日期统一为 YYYY-MM-DD，unit_nav 保留原始精度。
3. latest_unit_nav 只读取页面“最新单位净值/单位净值”旁明确印刷的数字；没有则填 null。
4. 不要根据曲线、收益率或前后行推算缺失值。
5. 少于 3 行可靠明细时仍原样输出已有行；没有则输出空数组。只输出 JSON。"""


def _recover_disclosed_nav_table(image_bytes: bytes) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Read exact dated unit NAV rows; never infer values from chart pixels."""
    from app.config import load_local_environment

    load_local_environment()
    provider_type = os.getenv("VLM_PROVIDER", "dashscope")
    model = os.getenv("VLM_MODEL", "qwen3-vl-flash")
    configured = bool(os.getenv("DASHSCOPE_API_KEY")) if provider_type == "dashscope" else bool(os.getenv("VLM_BASE_URL"))
    audit = {"configured": configured, "attempted": False, "succeeded": False,
             "provider": provider_type, "model": model, "calls": 0,
             "successful_calls": 0, "failed_calls": 0,
             "methods": ["VLM 披露净值表读取"], "stage": "disclosed_nav_table_recovery"}
    if not configured:
        audit["errors"] = ["VLM 未配置，无法读取披露净值表"]
        return [], audit
    try:
        from app.services.chart_extractor.vlm_extractor import create_provider

        provider = create_provider(provider_type, api_key=os.getenv("DASHSCOPE_API_KEY", ""),
                                   base_url=os.getenv("VLM_BASE_URL", ""), model=model)
        audit.update({"attempted": True})
        points: list[dict[str, Any]] = []
        latest_nav: float | None = None
        for variant_name, variant_bytes in _disclosed_nav_image_variants(image_bytes):
            audit["calls"] += 1
            raw = _run_async(provider.extract_structure(variant_bytes, _DISCLOSED_NAV_TABLE_PROMPT))
            audit["successful_calls"] += 1
            variant_points, variant_latest = _parse_disclosed_nav_payload(raw)
            if len(variant_points) > len(points):
                points, latest_nav = variant_points, variant_latest
                audit["selected_variant"] = variant_name
            if len(points) >= 3:
                break
        audit.update({"succeeded": True, "row_count": len(points)})
    except Exception as exc:
        audit.update({"attempted": True, "failed_calls": 1, "errors": [str(exc)[:300]]})
        return [], audit
    if len(points) < 3:
        audit.update({"succeeded": False, "errors": ["披露净值明细不足 3 行"]})
        return [], audit
    if latest_nav is not None and abs(float(points[-1]["nav"]) - latest_nav) / latest_nav > 0.005:
        audit.update({"succeeded": False, "errors": [
            f"历史表最新单位净值 {float(points[-1]['nav']):.6f} 与页面披露 {latest_nav:.6f} 不一致"]})
        return [], audit
    return points, audit


def _disclosed_nav_image_variants(image_bytes: bytes) -> list[tuple[str, bytes]]:
    """Retry tall screenshots on an enlarged lower section where history tables live."""
    variants = [("full_page", image_bytes)]
    try:
        from PIL import Image

        image = Image.open(BytesIO(image_bytes)).convert("RGB")
        width, height = image.size
        if height <= width * 1.2:
            return variants
        top = int(height * 0.68)
        lower = image.crop((0, top, width, height))
        scale = max(2, min(4, round(1800 / max(width, 1))))
        lower = lower.resize((lower.width * scale, lower.height * scale))
        output = BytesIO()
        lower.save(output, format="PNG")
        variants.append(("enlarged_lower_section", output.getvalue()))
    except Exception:
        pass
    return variants


def _parse_disclosed_nav_payload(raw: str) -> tuple[list[dict[str, Any]], float | None]:
    text = (raw or "").strip()
    if text.startswith("```"):
        text = "\n".join(line for line in text.splitlines() if not line.strip().startswith("```"))
    try:
        payload = json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{[\s\S]*\}", text)
        if not match:
            return [], None
        try:
            payload = json.loads(match.group())
        except json.JSONDecodeError:
            return [], None
    if not isinstance(payload, dict):
        return [], None
    points = _valid_nav_points([
        {"observation_date": item.get("date"), "nav": item.get("unit_nav")}
        for item in payload.get("nav_rows", []) if isinstance(item, dict)
    ])
    unique = {point["observation_date"]: point for point in points}
    points = [unique[key] for key in sorted(unique)]
    try:
        latest_nav = float(payload.get("latest_unit_nav")) if payload.get("latest_unit_nav") is not None else None
    except (TypeError, ValueError):
        latest_nav = None
    if latest_nav is not None and (not np.isfinite(latest_nav) or latest_nav <= 0 or latest_nav > 1_000_000):
        latest_nav = None
    return points, latest_nav


def _recover_monthly_nav_from_return_table(image_bytes: bytes) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Recover a reviewable low-frequency NAV candidate from a printed return table."""
    from app.config import load_local_environment

    load_local_environment()
    provider_type = os.getenv("VLM_PROVIDER", "dashscope")
    model = os.getenv("VLM_MODEL", "qwen3-vl-flash")
    configured = bool(os.getenv("DASHSCOPE_API_KEY")) if provider_type == "dashscope" else bool(os.getenv("VLM_BASE_URL"))
    audit = {
        "configured": configured,
        "attempted": False,
        "succeeded": False,
        "provider": provider_type,
        "model": model,
        "calls": 0,
        "successful_calls": 0,
        "failed_calls": 0,
        "methods": ["VLM 月收益表读取", "确定性复利重建"],
        "stage": "table_recovery",
    }
    if not configured:
        audit["errors"] = ["VLM 未配置，无法读取月收益表"]
        return [], audit
    try:
        from app.services.chart_extractor.vlm_extractor import create_provider

        provider = create_provider(
            provider_type,
            api_key=os.getenv("DASHSCOPE_API_KEY", ""),
            base_url=os.getenv("VLM_BASE_URL", ""),
            model=model,
        )
        audit.update({"attempted": True, "calls": 1})
        raw = _run_async(provider.extract_structure(image_bytes, _MONTHLY_RETURN_TABLE_PROMPT))
        rows, annual_returns, latest_nav = _parse_monthly_return_payload(raw)
        audit.update({"succeeded": True, "successful_calls": 1, "row_count": len(rows)})
    except Exception as exc:
        audit.update({"attempted": True, "failed_calls": 1, "errors": [str(exc)[:300]]})
        return [], audit
    if len(rows) < 2:
        audit.update({"succeeded": False, "errors": ["未读取到至少 2 个月的可靠收益单元格"]})
        return [], audit
    for year, reported_return in annual_returns.items():
        year_rows = [value for row_year, _, value in rows if row_year == year]
        if not year_rows:
            continue
        compounded = float(np.prod([1 + value for value in year_rows]) - 1)
        if abs(compounded - reported_return) > 0.012:
            audit.update({
                "succeeded": False,
                "failed_calls": 1,
                "errors": [f"{year}年月收益复利 {compounded:.2%} 与披露合计 {reported_return:.2%} 不一致"],
            })
            return [], audit
    first_year, first_month, _ = rows[0]
    baseline_year = first_year if first_month > 1 else first_year - 1
    baseline_month = first_month - 1 if first_month > 1 else 12
    baseline_day = calendar.monthrange(baseline_year, baseline_month)[1]
    points: list[dict[str, Any]] = [{
        "observation_date": date(baseline_year, baseline_month, baseline_day).isoformat(),
        "nav": 1.0,
    }]
    nav = 1.0
    for year, month, return_value in rows:
        nav *= 1 + return_value
        points.append({
            "observation_date": date(year, month, calendar.monthrange(year, month)[1]).isoformat(),
            "nav": round(nav, 8),
        })
    if latest_nav is not None and abs(nav - latest_nav) / latest_nav > 0.025:
        audit.update({
            "succeeded": False,
            "failed_calls": 1,
            "errors": [f"月收益重建终值 {nav:.4f} 与披露当前净值 {latest_nav:.4f} 不一致"],
        })
        return [], audit
    return _valid_nav_points(points), audit


def _parse_monthly_return_rows(raw: str) -> list[tuple[int, int, float]]:
    """Validate VLM table cells; reject future, duplicate and implausible rows."""
    rows, _, _ = _parse_monthly_return_payload(raw)
    return rows


def _parse_monthly_return_payload(
    raw: str,
) -> tuple[list[tuple[int, int, float]], dict[int, float], float | None]:
    text = (raw or "").strip()
    if text.startswith("```"):
        text = "\n".join(line for line in text.splitlines() if not line.strip().startswith("```"))
    try:
        payload = json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{[\s\S]*\}", text)
        if not match:
            return [], {}, None
        try:
            payload = json.loads(match.group())
        except json.JSONDecodeError:
            return [], {}, None
    raw_rows = payload.get("monthly_returns", []) if isinstance(payload, dict) else []
    unique: dict[tuple[int, int], float] = {}
    for item in raw_rows:
        if not isinstance(item, dict):
            continue
        try:
            year = int(item["year"])
            month = int(item["month"])
            return_value = float(item["return_pct"]) / 100
        except (KeyError, TypeError, ValueError):
            continue
        if not (1990 <= year <= date.today().year and 1 <= month <= 12 and -0.80 <= return_value <= 3.0):
            continue
        month_end = date(year, month, calendar.monthrange(year, month)[1])
        if month_end > date.today():
            continue
        key = (year, month)
        if key in unique and abs(unique[key] - return_value) > 1e-9:
            unique.pop(key, None)
            continue
        unique[key] = return_value
    annual_returns: dict[int, float] = {}
    for item in payload.get("annual_returns", []) if isinstance(payload, dict) else []:
        try:
            year = int(item["year"])
            value = float(item["return_pct"]) / 100
        except (KeyError, TypeError, ValueError):
            continue
        if 1990 <= year <= date.today().year and -0.95 <= value <= 10:
            annual_returns[year] = value
    try:
        latest_nav = float(payload.get("latest_nav")) if payload.get("latest_nav") is not None else None
    except (TypeError, ValueError):
        latest_nav = None
    if latest_nav is not None and (not np.isfinite(latest_nav) or latest_nav <= 0 or latest_nav > 1_000_000):
        latest_nav = None
    rows = [(year, month, unique[(year, month)]) for year, month in sorted(unique)]
    return rows, annual_returns, latest_nav


def _looks_like_strategy_label(value: str) -> bool:
    token = value.replace(" ", "")
    return bool(token) and any(marker in token for marker in ("策略", "产品介绍", "系列"))


def _normalise_manifest_product_name(value: str | None) -> str:
    """Turn a classified material label into a concrete product hint if possible."""
    token = (value or "").strip()
    token = re.sub(
        r"(?:近一年)?(?:产品)?(?:净值)?(?:业绩)?(?:周度报告|周报|月度报告|月报|运行报告)$",
        "",
        token,
    ).strip(" _-—（）()")
    generic = (
        "代表产品", "多产品", "产品总", "产品业绩", "策略代表产品", "期货多策略",
        "低波CTA", "中波CTA", "CTA策略", "CTA系列",
    )
    if len(token.replace(" ", "")) < 3 or any(label in token for label in generic):
        return ""
    if token.endswith("策略") and not re.search(r"\d+号|一期|二期|基金", token):
        return ""
    return token


def _is_multi_product_material(filename: str, product_name_hint: str | None) -> bool:
    token = f"{Path(filename).stem} {product_name_hint or ''}".replace(" ", "")
    return any(marker in token for marker in ("多产品", "产品总", "产品汇总", "产品组合", "产品池"))


def recover_pending_disclosed_nav(file_id: str) -> dict[str, Any]:
    """Run only the exact-table recovery lane for one unresolved image.

    The operation is idempotent: an existing candidate, including a reviewed
    or published one, prevents another machine version from being appended.
    """
    from app.services.upload_storage import resolve_upload_path

    with SessionFactory() as session:
        from app.models import RawFile

        raw_file = session.get(RawFile, file_id)
        if raw_file is None:
            return {"file_id": file_id, "outcome": "missing_file"}
        existing = session.execute(
            select(NavCandidateVersion.id).where(NavCandidateVersion.source_file_id == file_id)
        ).scalar()
        if existing is not None:
            return {"file_id": file_id, "outcome": "candidate_already_exists"}
        context = raw_file.ingestion_context or {}
        if _is_multi_product_material(raw_file.filename, context.get("product_name_hint")):
            audit = dict(raw_file.extraction_audit or {})
            workflow = dict(audit.get("workflow") or {})
            workflow.update({"binding": "needs_review", "trace": "blocked_on_binding", "next_action": "confirm_product_binding"})
            audit["workflow"] = workflow
            store.update_file_extraction_audit(session, file_id, audit)
            return {"file_id": file_id, "outcome": "needs_product_binding"}
        suffix = Path(raw_file.filename).suffix.lower()
        if suffix not in {".jpg", ".jpeg", ".png", ".webp"}:
            return {"file_id": file_id, "outcome": "unsupported_media"}
        path = resolve_upload_path(raw_file.storage_path)
        if path is None or not path.is_file():
            return {"file_id": file_id, "outcome": "missing_upload"}

        products = session.execute(
            select(ProductEntity).join(DocumentFragment, DocumentFragment.product_id == ProductEntity.id)
            .where(DocumentFragment.file_id == file_id).distinct()
        ).scalars().all()
        product = products[0] if len(products) == 1 else None
        if product is None:
            hint = _normalise_manifest_product_name(context.get("product_name_hint"))
            if hint:
                product = session.execute(
                    select(ProductEntity).where(ProductEntity.standard_name == hint)
                ).scalars().first()
        if product is None:
            audit = dict(raw_file.extraction_audit or {})
            workflow = dict(audit.get("workflow") or {})
            workflow.update({"identity": "needs_review", "next_action": "confirm_product_binding"})
            audit["workflow"] = workflow
            store.update_file_extraction_audit(session, file_id, audit)
            return {"file_id": file_id, "outcome": "needs_product_binding"}

        points, recovery_audit = _recover_disclosed_nav_table(path.read_bytes())
        audit = _merge_extraction_audit(raw_file.extraction_audit or {}, recovery_audit)
        if len(points) < 3:
            workflow = dict((raw_file.extraction_audit or {}).get("workflow") or {})
            workflow.update({"trace": "needs_calibration", "next_action": "recover_chart_calibration"})
            audit["workflow"] = workflow
            store.update_file_extraction_audit(session, file_id, audit)
            return {"file_id": file_id, "outcome": "no_reliable_table", "rows": len(points)}

        fragment = store.add_fragment(
            session,
            file_id=file_id,
            fragment_type="table",
            content_text=f"披露表格恢复 {len(points)} 个候选净值点",
            content_data={
                "method": "VLM 披露净值表读取",
                "review_status": "pending",
                "workflow_stage": "table_recovery",
                "num_points": len(points),
                "selected_variant": recovery_audit.get("selected_variant"),
            },
            ocr_confidence=0.9,
            product_id=product.id,
        )
        candidate = store.create_nav_candidate_version(
            session,
            product.id,
            file_id,
            points,
            source_fragment_id=fragment.id,
            frequency=_infer_frequency([date.fromisoformat(point["observation_date"]) for point in points]),
            confidence=0.9,
        )
        fragment.content_data = {**(fragment.content_data or {}), "candidate_version_id": candidate.id}
        session.commit()
        audit["workflow"] = {
            **dict((raw_file.extraction_audit or {}).get("workflow") or {}),
            "identity": "resolved",
            "trace": "candidate_created",
            "next_action": "review_candidate",
        }
        store.update_file_extraction_audit(session, file_id, audit)
        store.update_file_status(session, file_id, "completed")
        return {
            "file_id": file_id,
            "outcome": "candidate_created",
            "candidate_id": candidate.id,
            "product_name": product.standard_name,
            "rows": len(points),
        }


def _product_name_from_filename(filename: str) -> str:
    """Use a concrete filename product segment as the primary identity hint.

    Classified source images are commonly named ``管理人_产品名.png``.  This is
    more reliable than chart OCR, which regularly confuses axis labels such
    as “单位净值” for a product title.  Strategy-only files deliberately
    remain unbound and need a human to choose the product.
    """
    stem = Path(filename).stem.strip()
    if not stem:
        return ""
    candidate = stem.rsplit("_", 1)[-1].strip() if "_" in stem else stem
    return candidate if _looks_like_product_name(candidate) else ""


def _manager_name_from_filename(filename: str) -> str | None:
    stem = Path(filename).stem.strip()
    if "_" not in stem:
        return None
    manager = stem.rsplit("_", 1)[0].strip()
    return manager or None


def _looks_like_product_name(value: str) -> bool:
    token = value.replace(" ", "").strip()
    if len(token) < 3 or _looks_like_strategy_label(token):
        return False
    generic_labels = ("单位", "累计", "周度", "产品", "净值", "期末", "本期", "合计")
    return not any(label in token for label in generic_labels)


# ---------------------------------------------------------------------------
# Curve tracing (chart_extractor integration)
# ---------------------------------------------------------------------------


def _run_async(coro):
    """Bridge async chart_extractor calls from the sync background task."""
    try:
        asyncio.get_running_loop()
    except RuntimeError:
        return asyncio.run(coro)
    import concurrent.futures
    with concurrent.futures.ThreadPoolExecutor(1) as pool:
        return pool.submit(asyncio.run, coro).result()


def _trace_image_curves(
    session: Any,
    file_id: str,
    content: bytes,
    candidates: list,
    ordered_products: list,
    filename: str,
    metrics: list | None = None,
) -> list[dict[str, Any]]:
    """Crop each curve candidate and trace it into NAV observations."""
    import cv2

    img = cv2.imdecode(np.frombuffer(content, dtype=np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        return []
    h, w = img.shape[:2]
    stem = Path(filename).stem if filename else "未命名产品"

    metrics_by_id = {m.product_id: m for m in (metrics or [])}
    products_by_source_id = {
        metric.product_id: ordered_products[index]
        for index, metric in enumerate(metrics or [])
        if index < len(ordered_products) and ordered_products[index] is not None
    }
    single_product = ordered_products[0] if len(ordered_products) == 1 else None
    audits: list[dict[str, Any]] = []

    for idx, candidate in enumerate(candidates):
        x0, y0 = int(candidate.left_ratio * w), int(candidate.top_ratio * h)
        x1, y1 = int(candidate.right_ratio * w), int(candidate.bottom_ratio * h)
        crop = img[max(0, y0):min(h, y1), max(0, x0):min(w, x1)]
        if crop.size == 0:
            continue
        ok, buf = cv2.imencode(".png", crop)
        if not ok:
            continue

        # A VLM-provided source ID is authoritative. Positional matching is
        # only safe for a one-product report; for multiple products create a
        # pending candidate instead of attaching a curve to the wrong fund.
        candidate_product_id = getattr(candidate, "product_id", None)
        product = products_by_source_id.get(candidate_product_id) if candidate_product_id else None
        if product is None:
            product = single_product
        is_unmatched = product is None
        fallback_name = product.standard_name if product else f"{stem} / 待绑定曲线 {idx + 1}"
        # Keep the curve as unbound evidence rather than fabricating a
        # “待绑定曲线” product.  The user can bind it to a real product from
        # the source-file queue, and the formal product library stays clean.

        metric = metrics_by_id.get(candidate.product_id or "")
        disclosed = None
        if metric is not None:
            disclosed = {
                "start_date": metric.start_date,
                "end_date": metric.end_date,
                "cumulative_return": metric.cumulative_return,
            }

        points, frequency_hint, trace_confidence, audit = _extract_points_from_image(
            buf.tobytes(), fallback_name, disclosed, color_hint=getattr(candidate, "color_hex", None)
        )
        audits.append(audit)
        points = _valid_nav_points(points)
        if len(points) >= 3:
            if product is None:
                continue
            target = product

            # Cross-validate against existing table-sourced NAV data
            points = _cross_validate_with_table(session, target.id, points)

            binding_status = "unmatched" if is_unmatched else "matched"
            fragment = store.add_fragment(
                session,
                file_id=file_id,
                fragment_type="chart_traced",
                content_text=f"曲线追踪 {len(points)} 个数据点",
                content_data={
                    "curve_index": candidate.curve_index,
                    "num_points": len(points),
                    "method": _visual_method(),
                    "confidence": trace_confidence,
                    "review_status": "pending",
                    "binding_status": binding_status,
                    "binding_reason": None if binding_status == "matched" else "多产品曲线未返回可验证的产品 ID",
                    "binding_confidence": candidate.binding_confidence,
                    "binding_evidence": candidate.binding_evidence,
                    "legend_label": candidate.legend_label,
                    "layout_product_name": candidate.layout_product_name,
                    "color_hex": candidate.color_hex,
                },
                ocr_confidence=trace_confidence,
                product_id=target.id,
            )
            frequency = frequency_hint or _infer_frequency([
                    date.fromisoformat(point["observation_date"])
                    for point in points
                ], pixel_trace=True)
            candidate_version = store.create_nav_candidate_version(
                session, target.id, file_id, points,
                source_fragment_id=fragment.id, frequency=frequency, confidence=trace_confidence,
            )
            fragment.content_data = {**(fragment.content_data or {}), "candidate_version_id": candidate_version.id}
            session.commit()
            logger.info("Created candidate NAV version %s with %d points for %s from %s", candidate_version.id, len(points), target.standard_name, filename)
        else:
            # Never hide a failed trace behind a generic “completed” status.
            # Preserve the chart bounds and binding evidence so the browser
            # can open the exact source region for colour picking/calibration.
            store.add_fragment(
                session,
                file_id=file_id,
                fragment_type="chart_traced",
                content_text=f"曲线追踪未生成足够净值点（{len(points)} 个）",
                content_data={
                    "curve_index": candidate.curve_index,
                    "num_points": 0,
                    "method": _visual_method(),
                    "review_status": "pending",
                    "binding_status": "unmatched" if is_unmatched else "matched",
                    "binding_reason": "多产品曲线未返回可验证的产品 ID" if is_unmatched else "曲线像素追踪不足 3 点；需要重新取色或校准坐标轴",
                    "binding_confidence": candidate.binding_confidence,
                    "binding_evidence": candidate.binding_evidence,
                    "legend_label": candidate.legend_label,
                    "layout_product_name": candidate.layout_product_name,
                    "color_hex": candidate.color_hex,
                },
                ocr_confidence=0.2,
                product_id=product.id if product else None,
            )

    return audits


def _trace_whole_image(session: Any, file_id: str, content: bytes, product_name: str) -> dict[str, Any]:
    """Trace the entire image as a single NAV chart (no table/candidates found)."""
    points, frequency_hint, trace_confidence, audit = _extract_points_from_image(content, product_name)
    points = _valid_nav_points(points)
    if len(points) < 3:
        return audit
    product = _find_or_create_product(session, product_name, None, None, None)
    fragment = store.add_fragment(
        session,
        file_id=file_id,
        fragment_type="chart_traced",
        content_text=f"整图曲线追踪 {len(points)} 个数据点",
        content_data={"num_points": len(points), "method": _visual_method(), "confidence": trace_confidence, "review_status": "pending"},
        ocr_confidence=trace_confidence,
        product_id=product.id,
    )
    frequency = frequency_hint or _infer_frequency([
            date.fromisoformat(point["observation_date"])
            for point in points
        ], pixel_trace=True)
    candidate_version = store.create_nav_candidate_version(
        session, product.id, file_id, points,
        source_fragment_id=fragment.id, frequency=frequency, confidence=trace_confidence,
    )
    fragment.content_data = {**(fragment.content_data or {}), "candidate_version_id": candidate_version.id}
    session.commit()
    return audit


def _cross_validate_with_table(
    session: Any,
    product_id: str,
    points: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Correct chart-traced NAV values using existing table-sourced observations.

    If the product already has NAV data from a table (higher confidence),
    find matching dates and apply a linear correction (scale + offset) to
    fix systematic Y-axis calibration errors in the chart tracing.
    """
    from sqlalchemy import select as sa_select

    from app.models import NavObservation

    # Get existing observations for this product (table-sourced = ground truth)
    existing = session.execute(
        sa_select(NavObservation)
        .where(NavObservation.product_id == product_id)
        .order_by(NavObservation.observation_date)
    ).scalars().all()

    if len(existing) < 2:
        return points

    # Build reference map: date_str → nav
    ref_map: dict[str, float] = {}
    for obs in existing:
        ref_map[obs.observation_date.isoformat()] = obs.nav

    # Find matching points
    matched_traced: list[float] = []
    matched_ref: list[float] = []
    for p in points:
        d = p.get("observation_date", "")
        v = p.get("nav")
        if d and v and d in ref_map:
            matched_traced.append(float(v))
            matched_ref.append(ref_map[d])

    if len(matched_traced) < 2:
        return points

    # Least-squares linear fit: ref = scale * traced + offset
    n = len(matched_traced)
    sum_x = sum(matched_traced)
    sum_y = sum(matched_ref)
    sum_xy = sum(x * y for x, y in zip(matched_traced, matched_ref))
    sum_x2 = sum(x * x for x in matched_traced)

    denom = n * sum_x2 - sum_x * sum_x
    if abs(denom) < 1e-12:
        offset = sum_y / n - sum_x / n
        scale = 1.0
    else:
        scale = (n * sum_xy - sum_x * sum_y) / denom
        offset = (sum_y - scale * sum_x) / n

    # Sanity check
    if scale < 0.5 or scale > 2.0:
        logger.warning("Cross-validation scale=%.3f out of range, skipping", scale)
        return points

    # Apply correction
    for p in points:
        v = p.get("nav")
        if v:
            corrected = scale * float(v) + offset
            if corrected > 0:
                p["nav"] = round(corrected, 6)

    logger.info("Table cross-validation: %d matches, scale=%.4f, offset=%.6f", n, scale, offset)
    return points


def _extract_points_from_image(
    image_bytes: bytes,
    curve_name: str,
    disclosed: dict[str, Any] | None = None,
    color_hint: str | None = None,
) -> tuple[list[dict[str, Any]], str | None, float, dict[str, Any]]:
    """Run chart extraction on one image; return NAV point dicts.

    When a VLM is configured it reads the chart structure (and axis anchors)
    and the standard pipeline calibrates the trace. In CV-only mode there are
    no axis anchors, so the curve is traced as raw pixels and calibrated
    against the disclosed table metrics (start/end dates + cumulative return)
    via endpoint anchoring.
    """
    # DashScope needs an API key; local / OpenAI-compatible vLLM only needs a
    # reachable base URL.  Treat both as a configured visual fallback.
    from app.config import load_local_environment

    load_local_environment()
    provider_type = os.getenv("VLM_PROVIDER", "dashscope")
    use_vlm = bool(os.getenv("DASHSCOPE_API_KEY", "")) if provider_type == "dashscope" else bool(os.getenv("VLM_BASE_URL", ""))

    if not use_vlm:
        points = _extract_points_cv_only(image_bytes, disclosed)
        return points, None, 0.55 if len(points) >= 3 else 0.2, {
            "configured": False,
            "attempted": False,
            "succeeded": False,
            "methods": ["CV 像素追踪（无 VLM）"],
        }

    from app.services.chart_extractor import extract_single_image

    def points_from_result(result: Any) -> list[dict[str, Any]]:
        points: list[dict[str, Any]] = []
        for curve in result.curves:
            if curve.is_benchmark:
                continue
            raw = [p for p in curve.points if p.value is not None and p.value > 0]
            for p in curve.points:
                if p.date and p.value is not None and p.value > 0:
                    points.append({"observation_date": p.date, "nav": p.value})
            if points:
                break  # one product curve per crop
            # A VLM often reads the frame, colour and Y scale correctly but
            # omits dense Chinese X-axis labels.  When the same factsheet
            # discloses a complete start/end period, map the already traced
            # pixel sequence across that interval and retain weekly samples.
            # This is explicitly interpolation evidence, never OCR dates.
            if len(raw) >= 3 and disclosed and disclosed.get("start_date") and disclosed.get("end_date"):
                try:
                    start = date.fromisoformat(str(disclosed["start_date"]))
                    end = date.fromisoformat(str(disclosed["end_date"]))
                    span = max((end - start).days, 1)
                    left, right = raw[0].x_px, raw[-1].x_px
                    width = max(right - left, 1)
                    points = [
                        {
                            "observation_date": (start.fromordinal(start.toordinal() + round((point.x_px - left) / width * span))).isoformat(),
                            "nav": float(point.value),
                        }
                        for point in raw
                    ]
                except (TypeError, ValueError):
                    points = []
            if points:
                break
        return points

    # Do not pass an auto-detected colour when VLM is enabled: a non-empty
    # curve_specs list intentionally switches the chart pipeline into manual
    # mode and silently skips VLM.  Let the model identify the line first, then
    # use CV colour tracing for the numerical pixels.  A manual CV pass remains
    # the deterministic fallback when the model is unavailable or uncertain.
    if use_vlm:
        audit = {
            "configured": True,
            "attempted": False,
            "succeeded": False,
            "provider": provider_type,
            "model": os.getenv("VLM_MODEL", "qwen3-vl-flash"),
            "calls": 0,
            "successful_calls": 0,
            "failed_calls": 0,
            "errors": [],
            "methods": ["VLM 结构识别", "CV 像素追踪"],
        }
        try:
            curve_specs = None
            if color_hint:
                curve_specs = [{
                    "name": curve_name,
                    "color_hex": color_hint,
                    "color_name": "layout_hint",
                    "is_benchmark": False,
                }]
            result = _run_async(extract_single_image(image_bytes, curve_specs=curve_specs, use_vlm=True))
            audit.update({
                "attempted": bool(result.vlm_attempted),
                "succeeded": bool(result.vlm_succeeded),
                "provider": result.vlm_provider or provider_type,
                "model": result.vlm_model or os.getenv("VLM_MODEL", "qwen3-vl-flash"),
                "calls": 1 if result.vlm_attempted else 0,
                "successful_calls": 1 if result.vlm_succeeded else 0,
                "failed_calls": 1 if result.vlm_attempted and not result.vlm_succeeded else 0,
                "errors": [result.vlm_error] if result.vlm_error else [],
            })
            points = points_from_result(result)
            # A frequency is authoritative only when the VLM also returned
            # dated observations.  On a pixel-only fallback, a model may say
            # "daily" simply because the chart is dense; those are still
            # interpolation samples and are handled as weekly below.
            frequency_hint = result.frequency if result.frequency in {"daily", "weekly", "monthly", "quarterly"} and points else None
            if len(points) >= 3:
                return _resample_pixel_points_weekly(points), frequency_hint, max(0.3, min(0.98, float(result.confidence))), audit
            logger.info("VLM chart structure returned no usable points for %s; falling back to CV", curve_name)
        except Exception as exc:
            logger.warning("VLM chart extraction failed for %s: %s", curve_name, exc)
            frequency_hint = None
            audit.update({"attempted": True, "calls": 1, "failed_calls": 1, "errors": [str(exc)[:300]]})
    else:
        frequency_hint = None

    # VLM often identifies the line and Y-axis but cannot read every X-axis
    # date on a dense report.  In that case its points have empty dates and
    # cannot enter the NAV table.  Use the deterministic CV tracer with the
    # disclosed endpoints to generate dated values instead of sending an
    # uncalibrated manual extraction through the VLM pipeline again.
    cv_points = _extract_points_cv_only(image_bytes, disclosed)
    if len(cv_points) >= 3:
        return cv_points, frequency_hint, 0.55, audit

    # Last resort for a chart with no disclosed period: a manual colour trace
    # may still produce reviewable pixel evidence, even though it has no date
    # calibration and will be kept out of the research-ready NAV series.
    colors = _auto_curve_colors(image_bytes)
    if not colors:
        return [], frequency_hint, 0.2, audit
    curve_specs = [{"name": curve_name, "color_hex": colors[0], "color_name": "", "is_benchmark": False}]
    try:
        result = _run_async(extract_single_image(image_bytes, curve_specs=curve_specs, use_vlm=False))
    except Exception as exc:
        logger.warning("CV chart extraction failed: %s", exc)
        audit.update({"failed_calls": int(audit.get("failed_calls", 0)) + 1, "errors": [*audit.get("errors", []), str(exc)[:300]]})
        return [], frequency_hint, 0.2, audit
    points = points_from_result(result)
    return points, frequency_hint, max(0.3, min(0.8, float(getattr(result, "confidence", 0.55)))), audit


def _extract_points_cv_only(
    image_bytes: bytes,
    disclosed: dict[str, Any] | None,
    color_hint: str | None = None,
) -> list[dict[str, Any]]:
    """Trace a curve with pure CV and calibrate it to real NAV values.

    No VLM / axis labels are used here. The curve is traced as raw pixels and
    then anchored to the disclosed start/end NAV values (endpoint calibration):
    leftmost traced pixel -> (start_date, 1.0), rightmost -> (end_date,
    1 + cumulative_return). Dates are interpolated linearly between the two
    endpoints. Axis-label OCR was evaluated and rejected as too unreliable.
    """
    import cv2

    from app.services.chart_extractor import (
        calibrate_from_disclosed,
        sanitize_nav_points,
    )
    from app.services.chart_extractor.pixel_tracer import detect_plot_area, trace_curve

    colors = [color_hint] if color_hint else _auto_curve_colors(image_bytes)
    if not colors:
        return []

    img = cv2.imdecode(np.frombuffer(image_bytes, dtype=np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        return []

    plot_area = detect_plot_area(img)
    points = trace_curve(img, plot_area, color_hex=colors[0])
    if len(points) < 3:
        return []

    if disclosed:
        points = calibrate_from_disclosed(
            points,
            start_date=disclosed.get("start_date", ""),
            end_date=disclosed.get("end_date", ""),
            cumulative_return=float(disclosed.get("cumulative_return", 0.0)),
            start_nav=float(disclosed.get("start_nav", 1.0)),
        )

    points = sanitize_nav_points(points)
    raw_points = [
        {"observation_date": p.date, "nav": p.value}
        for p in points
        if p.date and p.value is not None and p.value > 0
    ]
    return _resample_pixel_points_weekly(raw_points)


def _resample_pixel_points_weekly(points: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Turn dense per-pixel samples into one observation every seven days.

    The CV tracer emits one point for almost every image column.  Those are
    interpolation samples, not observations published by the manager.  Taking
    the last available pixel at each seven-day target preserves the curve
    shape and endpoints while preventing anti-aliasing noise from becoming
    artificial daily returns.
    """
    if len(points) < 3:
        return points
    points = _normalise_nav_points(points)
    ordered = sorted(
        (
            item
            for item in points
            if item.get("observation_date") and item.get("nav") is not None
        ),
        key=lambda item: str(item["observation_date"]),
    )
    if len(ordered) < 3:
        return points
    dates = [date.fromisoformat(item["observation_date"]) for item in ordered]
    if (dates[-1] - dates[0]).days < 30:
        return ordered

    sampled: list[dict[str, Any]] = []
    cursor = dates[0]
    index = 0
    while cursor <= dates[-1]:
        while index + 1 < len(ordered) and dates[index + 1] <= cursor:
            index += 1
        source = ordered[index]
        sampled.append({"observation_date": cursor, "nav": source["nav"]})
        cursor = cursor.fromordinal(cursor.toordinal() + 7)

    # Always retain the disclosed end point; otherwise a one-week endpoint
    # mismatch can make the table's cumulative return appear inconsistent.
    if sampled[-1]["observation_date"] != dates[-1]:
        sampled.append({"observation_date": dates[-1], "nav": ordered[-1]["nav"]})
    return sampled


def _auto_curve_colors(image_bytes: bytes, max_colors: int = 2) -> list[str]:
    """Detect dominant curve colors via HSV scan (red/green/blue families)."""
    import cv2

    img = cv2.imdecode(np.frombuffer(image_bytes, dtype=np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        return []
    h, w = img.shape[:2]
    roi = img[int(h * 0.1):int(h * 0.85), int(w * 0.08):int(w * 0.95)]
    if roi.size == 0:
        return []
    hsv = cv2.cvtColor(roi, cv2.COLOR_BGR2HSV)
    hue, sat, val = hsv[:, :, 0], hsv[:, :, 1], hsv[:, :, 2]
    saturated = (sat >= 60) & (val >= 80)

    families: list[tuple[int, np.ndarray]] = []
    red_mask = saturated & ((hue <= 10) | (hue >= 170))
    green_mask = saturated & (hue >= 35) & (hue <= 90)
    blue_mask = saturated & (hue >= 100) & (hue <= 135)
    for mask in (red_mask, green_mask, blue_mask):
        count = int(mask.sum())
        if count > (roi.shape[0] * roi.shape[1]) * 0.001:
            families.append((count, mask))

    families.sort(key=lambda item: item[0], reverse=True)
    colors: list[str] = []
    for _, mask in families[:max_colors]:
        # Average only the masked (curve) pixels; averaging the bitwise_and
        # output would dilute the color with the zeroed background.
        mean_bgr = cv2.mean(roi, mask=mask.astype(np.uint8))
        b, g, r = int(mean_bgr[0]), int(mean_bgr[1]), int(mean_bgr[2])
        colors.append(f"#{r:02x}{g:02x}{b:02x}")
    return colors


def _ingest_nav_table(session: Any, file_id: str, content: bytes, filename: str, suffix: str) -> int:
    frame = pd.read_csv(BytesIO(content)) if suffix == ".csv" else pd.read_excel(BytesIO(content))
    return _ingest_nav_frame(session, file_id, frame, filename)


def _ingest_nav_frame(
    session: Any,
    file_id: str,
    frame: pd.DataFrame,
    filename: str,
    *,
    location: str | None = None,
) -> int:
    """Persist a disclosed NAV table only when its dates and values are usable."""
    date_column = _find_column(frame, ("date", "日期", "净值日期", "估值日期", "观察日"))
    if date_column is None:
        # Heuristic: first column mostly parseable as dates.
        for col in frame.columns:
            if pd.to_datetime(frame[col], errors="coerce", format="mixed").notna().mean() > 0.7:
                date_column = col
                break
    if date_column is None:
        return 0

    nav_column = _find_column(frame, ("nav", "净值", "单位净值", "累计净值", "acc_nav"))
    if nav_column is not None:
        # Long format: single explicit NAV column.
        value_columns = [nav_column]
    else:
        # Wide format: every numeric column (except date) is one product.
        value_columns = [
            col for col in frame.columns
            if col != date_column and pd.api.types.is_numeric_dtype(frame[col])
        ]
        if not value_columns:
            return 0

    dates_series = pd.to_datetime(frame[date_column], errors="coerce", format="mixed")
    total_rows = 0
    for col in value_columns:
        values = pd.to_numeric(frame[col], errors="coerce")
        mask = dates_series.notna() & values.notna() & (values > 0)
        parsed = pd.DataFrame({"date": dates_series[mask], "nav": values[mask]}).drop_duplicates("date").sort_values("date")
        if len(parsed) < 2:
            continue
        dates = [value.date() for value in parsed["date"].tolist()]
        col_name = str(col).strip()
        product_name = Path(filename).stem if nav_column is not None else col_name
        if not product_name or product_name.startswith("Unnamed"):
            product_name = Path(filename).stem
        product = _find_or_create_product(session, product_name, None, dates[0], dates[-1])
        fragment = store.add_fragment(
            session,
            file_id=file_id,
            fragment_type="table",
            content_text=f"自动导入 {len(parsed)} 条净值记录（列: {col_name}）",
            content_data={"date_column": str(date_column), "nav_column": col_name, "rows": len(parsed), "method": "表格解析（Pandas）", "review_status": "pending", "location": location},
            ocr_confidence=1.0,
            product_id=product.id,
        )
        frequency = _infer_frequency(dates)
        total_rows += store.add_nav_observations(
            session,
            product.id,
            [{"observation_date": item_date, "nav": float(value)} for item_date, value in zip(dates, parsed["nav"].tolist(), strict=True)],
            source_file_id=file_id,
            source_fragment_id=fragment.id,
            frequency=frequency,
        )
    if total_rows == 0:
        return 0
    return total_rows


def _ingest_office_document(session: Any, file_id: str, content: bytes, filename: str, suffix: str) -> int:
    """Extract Word/PPT text and disclosed tables without inventing chart data."""
    product = _find_or_create_product(session, Path(filename).stem, None, None, None)
    parts = _extract_office_parts(content, suffix)
    if not any(part["text"] or part["tables"] for part in parts):
        raise ValueError("文档未提取到可读取的文本或表格")

    total_rows = 0
    for part in parts:
        location = part["location"]
        if part["text"]:
            store.add_fragment(
                session,
                file_id=file_id,
                fragment_type="text",
                page_number=part["page_number"],
                content_text=part["text"][:20000],
                content_data={"location": location, "method": part["method"], "review_status": "pending"},
                ocr_confidence=1.0,
                product_id=product.id,
            )
        for table_index, table in enumerate(part["tables"], 1):
            if len(table) < 2:
                continue
            header, *rows = table
            if not header or not rows:
                continue
            header_tokens = [str(value).replace(" ", "").lower() for value in header]
            has_explicit_date = any(any(label in token for label in ("date", "日期", "净值日期", "估值日期")) for token in header_tokens)
            has_explicit_nav = any(any(label in token for label in ("nav", "净值", "单位净值", "累计净值", "acc_nav")) for token in header_tokens)
            if not (has_explicit_date and has_explicit_nav):
                continue
            frame = pd.DataFrame(rows, columns=header)
            total_rows += _ingest_nav_frame(
                session,
                file_id,
                frame,
                filename,
                location=f"{location}; table={table_index}",
            )
    return total_rows


def _extract_office_parts(content: bytes, suffix: str) -> list[dict[str, Any]]:
    """Return text/table evidence with stable document locations for Office files."""
    if suffix == ".docx":
        from docx import Document

        document = Document(BytesIO(content))
        text = "\n".join(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
        tables = [[ [cell.text.strip() for cell in row.cells] for row in table.rows ] for table in document.tables]
        return [{"page_number": None, "location": "Word document", "text": text, "tables": tables, "method": "DOCX XML 提取"}]

    parts: list[dict[str, Any]] = []
    # Read slide XML directly. python-pptx materializes embedded media and can
    # consume gigabytes on image-heavy manager decks even though ingestion only
    # needs text/table evidence.
    from lxml import etree

    with zipfile.ZipFile(BytesIO(content)) as archive:
        slide_names = sorted(
            (
                name for name in archive.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            ),
            key=lambda name: int(re.search(r"(\d+)\.xml$", name).group(1)),
        )
        for index, slide_name in enumerate(slide_names, 1):
            root = etree.fromstring(archive.read(slide_name))
            namespaces = {
                "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
            }
            text_values = [value.strip() for value in root.xpath("//a:t/text()", namespaces=namespaces) if value.strip()]
            tables: list[list[list[str]]] = []
            for table in root.xpath("//a:tbl", namespaces=namespaces):
                rows: list[list[str]] = []
                for row in table.xpath("./a:tr", namespaces=namespaces):
                    cells = ["".join(cell.xpath(".//a:t/text()", namespaces=namespaces)).strip() for cell in row.xpath("./a:tc", namespaces=namespaces)]
                    rows.append(cells)
                if rows:
                    tables.append(rows)
            parts.append({"page_number": index, "location": f"slide={index}", "text": "\n".join(text_values), "tables": tables, "method": "PPTX XML 流式提取"})
    return parts


def _ingest_pdf(
    session: Any,
    file_id: str,
    content: bytes,
    filename: str,
    product_name_hint: str | None = None,
) -> dict[str, Any]:
    stem = Path(filename).stem

    # Phase 1: text layer for context.
    text = ""
    try:
        from pypdf import PdfReader
        text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages).strip()
    except Exception as exc:
        logger.info("PDF text extraction skipped: %s", exc)

    if text:
        store.add_fragment(
            session,
            file_id=file_id,
            fragment_type="text",
            content_text=text[:20000],
            ocr_confidence=0.9,
            product_id=None,
        )

    # Phase 2: chart pipeline for NAV curves.  Text-only PDFs still continue
    # to the deterministic monthly-table lane below; a missing curve color is
    # not evidence that the PDF belongs to the filename's product.
    provider_type = os.getenv("VLM_PROVIDER", "dashscope")
    use_vlm = bool(os.getenv("DASHSCOPE_API_KEY", "")) if provider_type == "dashscope" else bool(os.getenv("VLM_BASE_URL", ""))
    curve_specs = None
    audits: list[dict[str, Any]] = []
    chart_job = None
    if not use_vlm:
        colors = _auto_curve_colors_from_pdf(content)
        if colors:
            curve_specs = [{"name": stem, "color_hex": colors[0], "color_name": "", "is_benchmark": False}]
        elif not text:
            raise ValueError("PDF 无可提取文本且未检测到净值曲线颜色，请在研究工具中手动校准")

    if use_vlm or curve_specs:
        from app.services.chart_extractor import run_extraction

        try:
            chart_job = _run_async(run_extraction(
                pdf_bytes=content,
                filename=filename,
                use_vlm=use_vlm,
                curve_specs=curve_specs,
            ))
        except Exception as exc:
            logger.warning("PDF chart extraction failed: %s", exc)
            if not text:
                raise ValueError(f"PDF 图表提取失败：{exc}") from exc
            audits.append({
                "configured": use_vlm,
                "attempted": use_vlm,
                "succeeded": False,
                "provider": provider_type if use_vlm else None,
                "model": os.getenv("VLM_MODEL") if use_vlm else None,
                "calls": 0,
                "successful_calls": 0,
                "failed_calls": 0,
                "errors": [str(exc)[:300]],
                "methods": ["PDF 文本层提取"],
            })

    traced_products = 0
    traced_targets: list[ProductEntity] = []
    if chart_job is not None:
        for result_index, result in enumerate(chart_job.results):
            audits.append({
                "configured": use_vlm,
                "attempted": bool(result.vlm_attempted),
                "succeeded": bool(result.vlm_succeeded),
                "provider": result.vlm_provider or provider_type if use_vlm else None,
                "model": result.vlm_model or os.getenv("VLM_MODEL", "qwen3-vl-flash") if use_vlm else None,
                "calls": 1 if result.vlm_attempted else 0,
                "successful_calls": 1 if result.vlm_succeeded else 0,
                "failed_calls": 1 if result.vlm_attempted and not result.vlm_succeeded else 0,
                "errors": [result.vlm_error] if result.vlm_error else [],
                "methods": ["VLM 结构识别", "CV 像素追踪"] if use_vlm else ["CV 像素追踪（无 VLM）"],
            })
            source_page = chart_job.regions[result_index].page_index if result_index < len(chart_job.regions) else 0
            for curve in result.curves:
                if curve.is_benchmark:
                    continue
                points = [
                    {"observation_date": p.date, "nav": p.value}
                    for p in curve.points
                    if p.date and p.value is not None and p.value > 0
                ]
                if len(points) < 3:
                    continue
                curve_name = (curve.name or "").strip()
                title = (result.structure.chart_title.strip() if result.structure and result.structure.chart_title else "")
                target_name = curve_name or title or stem
                if not _looks_like_pdf_product_curve(target_name, stem):
                    logger.info("Skipping non-product PDF curve label %r from %s", target_name, filename)
                    continue
                target = _find_or_create_product(session, target_name, None, None, None)
                if target not in traced_targets:
                    traced_targets.append(target)
                fragment = store.add_fragment(
                    session,
                    file_id=file_id,
                    fragment_type="chart_traced",
                    page_number=source_page + 1,
                    content_text=f"PDF 曲线追踪 {len(points)} 个数据点",
                    content_data={"num_points": len(points), "confidence": result.confidence, "source_page": source_page + 1, "method": "VLM 结构识别 + CV 像素追踪" if use_vlm else "CV 像素追踪（无 VLM）", "review_status": "pending"},
                    ocr_confidence=max(0.3, result.confidence * 0.8),
                    product_id=target.id,
                )
                frequency = result.frequency if result.frequency != "unknown" else None
                candidate_version = store.create_nav_candidate_version(
                    session, target.id, file_id, points,
                    source_fragment_id=fragment.id,
                    frequency=frequency,
                    confidence=max(0.3, result.confidence * 0.8),
                )
                fragment.content_data = {**(fragment.content_data or {}), "candidate_version_id": candidate_version.id}
                session.commit()
                traced_products += 1

    reliable_curve = session.execute(
        select(NavCandidateVersion.id).where(
            NavCandidateVersion.source_file_id == file_id,
            NavCandidateVersion.confidence >= 0.65,
        )
    ).scalar() is not None
    monthly_points = _recover_monthly_nav_from_pdf_text(text, filename=filename) if not reliable_curve else []
    recovery_target = traced_targets[0] if len(traced_targets) == 1 else None
    if recovery_target is None and not traced_targets:
        hinted_name = _normalise_manifest_product_name(product_name_hint)
        if hinted_name:
            recovery_target = _find_or_create_product(session, hinted_name, None, None, None)
    if len(monthly_points) >= 3 and recovery_target is not None:
        target = recovery_target
        fragment = store.add_fragment(
            session,
            file_id=file_id,
            fragment_type="table",
            page_number=1,
            content_text=f"PDF 文本层月收益表恢复 {len(monthly_points)} 个候选净值点",
            content_data={
                "method": "PDF 文本层月收益表 + 确定性复利重建",
                "review_status": "pending",
                "workflow_stage": "table_recovery",
                "num_points": len(monthly_points),
            },
            ocr_confidence=0.92,
            product_id=target.id,
        )
        candidate = store.create_nav_candidate_version(
            session,
            target.id,
            file_id,
            monthly_points,
            source_fragment_id=fragment.id,
            frequency="monthly",
            confidence=0.92,
        )
        fragment.content_data = {**(fragment.content_data or {}), "candidate_version_id": candidate.id}
        session.commit()
    elif len(monthly_points) >= 3:
        audits.append({
            "configured": False,
            "attempted": False,
            "succeeded": False,
            "errors": ["PDF 月收益表已读取，但缺少唯一产品绑定；未创建净值候选"],
            "methods": ["PDF 文本层月收益表（待产品绑定）"],
        })

    if traced_products == 0 and not text:
        raise ValueError("PDF 未提取到文本或净值曲线，请在研究工具中手动校准")
    return _merge_extraction_audit(*audits)


def _recover_monthly_nav_from_pdf_text(text: str, filename: str | None = None) -> list[dict[str, Any]]:
    """Rebuild a reviewable NAV candidate from a PDF monthly-return table.

    PDF text extraction often emits the table before the title/metadata and
    may omit the header entirely.  The report period therefore comes from the
    title when available, then from a ``YYMM月报`` filename token.  When the
    share inception date is not present before the table, the first visible
    monthly cell is used and the preceding month-end becomes the baseline.
    """
    if not text or "月收益" not in text:
        # Some CJK PDF fonts are not decoded by pypdf.  The rows themselves
        # are still useful when a report-month token is available in the
        # filename, so do not require the literal section heading in that
        # case.
        if not filename or not re.search(r"月(?:报|度)", Path(filename).stem):
            return []
    report_match = re.search(r"(20\d{2})年\s*(\d{1,2})月", text)
    if report_match:
        report_year, report_month = int(report_match.group(1)), int(report_match.group(2))
    else:
        report_year = report_month = 0
        if filename:
            filename_match = re.search(
                r"(?<!\d)(?:(20)?(\d{2}))(0[1-9]|1[0-2])月(?:报|度)",
                Path(filename).stem,
            )
            if filename_match:
                century, short_year, month = filename_match.groups()
                report_year = int(f"{century or '20'}{short_year}")
                report_month = int(month)
    if not report_year or not report_month:
        return []

    # Dates after the monthly table are usually chart x-axis labels.  Only
    # trust a full date occurring in the metadata prefix; otherwise infer a
    # month-end baseline from the first monthly row below.
    table_start = text.find("月收益")
    metadata_text = text[:table_start] if table_start >= 0 else ""
    inception_match = re.search(r"(20\d{2})[/-](\d{1,2})[/-](\d{1,2})", metadata_text)
    nav_match = re.search(r"(?:累计净值|单位净值)\s+([0-9]+(?:\.[0-9]+)?)", text)

    # Keep the longest percentage row per year.  The same PDF commonly has a
    # second quarterly table; selecting the longest row prevents those
    # summary rows from being mistaken for monthly data.
    raw_rows: dict[int, list[float]] = {}
    for line in text.splitlines():
        match = re.match(r"^\s*(20\d{2})\s+(.+)$", line.strip())
        if not match:
            continue
        year = int(match.group(1))
        values = [float(value.rstrip("%")) / 100 for value in re.findall(r"[-+]?\d+(?:\.\d+)?%", match.group(2))]
        if not values or year > report_year:
            continue
        if len(values) > len(raw_rows.get(year, [])):
            raw_rows[year] = values
    if not raw_rows:
        return []

    explicit_inception = None
    if inception_match:
        explicit_inception = date(
            int(inception_match.group(1)),
            int(inception_match.group(2)),
            int(inception_match.group(3)),
        )
    first_year = min(raw_rows)
    if explicit_inception is not None:
        inception = explicit_inception
        first_start_month = inception.month + 1
        first_expected_months = 12 - inception.month
    else:
        # When the text layer loses CJK labels, chart x-axis dates remain
        # machine-readable.  They give an unambiguous first monthly cell and
        # avoid mistaking December's return for an annual-total column.
        chart_dates = [
            date(int(year), int(month), int(day))
            for year, month, day in re.findall(r"(20\d{2})[/-](\d{1,2})[/-](\d{1,2})", text)
        ]
        first_chart_date = next((item for item in sorted(chart_dates) if item.year == first_year), None)
        if first_chart_date is not None:
            first_start_month = first_chart_date.month
            first_expected_months = 13 - first_start_month
            inception = date(
                first_chart_date.year,
                first_chart_date.month - 1 if first_chart_date.month > 1 else 12,
                calendar.monthrange(
                    first_chart_date.year if first_chart_date.month > 1 else first_chart_date.year - 1,
                    first_chart_date.month - 1 if first_chart_date.month > 1 else 12,
                )[1],
            )
        else:
        # Prefer treating the final cell as an annual total only when it
        # independently matches the preceding cells' compounded return.
            first_values = raw_rows[first_year]
            annual_candidate = first_values[-1] if len(first_values) >= 2 else None
            month_candidate = first_values[:-1] if annual_candidate is not None else []
            annual_is_proven = bool(
                month_candidate
                and annual_candidate is not None
                and abs(float(np.prod([1 + value for value in month_candidate]) - 1) - annual_candidate) <= 0.012
            )
            month_count = len(month_candidate) if annual_is_proven else len(first_values)
            if month_count < 1 or month_count > 12:
                return []
            first_start_month = 13 - month_count
            baseline_year = first_year if first_start_month > 1 else first_year - 1
            baseline_month = first_start_month - 1 if first_start_month > 1 else 12
            inception = date(baseline_year, baseline_month, calendar.monthrange(baseline_year, baseline_month)[1])
            first_expected_months = month_count

    rows: list[tuple[int, list[float], float | None]] = []
    for year in sorted(raw_rows):
        values = list(raw_rows[year])
        if year < inception.year or year > report_year:
            continue
        expected_months = (
            first_expected_months
            if year == inception.year
            else report_month
            if year == report_year
            else 12
        )
        annual: float | None = None
        if len(values) == expected_months + 1:
            annual = values.pop()
        elif year != first_year and 2 <= len(values) <= expected_months + 1:
            # Some PDF text layers omit visually empty month cells.  A short
            # row is usable only when its final printed value independently
            # proves that the preceding cells are monthly returns.
            possible_annual = values[-1]
            possible_months = values[:-1]
            compounded = float(np.prod([1 + value for value in possible_months]) - 1)
            if abs(compounded - possible_annual) <= 0.012:
                values = possible_months
                annual = possible_annual
        if not values or len(values) > expected_months:
            continue
        if annual is not None:
            compounded = float(np.prod([1 + value for value in values]) - 1)
            if abs(compounded - annual) > 0.012:
                return []
        rows.append((year, values, annual))
    if not rows or rows[0][0] != first_year or rows[-1][0] != report_year:
        return []
    points: list[dict[str, Any]] = [{
        "observation_date": inception.isoformat(),
        "nav": 1.0,
    }]
    nav = 1.0
    for year, values, _ in rows:
        start_month = first_start_month if year == inception.year else 1
        for offset, return_value in enumerate(values):
            month = start_month + offset
            nav *= 1 + return_value
            points.append({
                "observation_date": date(year, month, calendar.monthrange(year, month)[1]).isoformat(),
                "nav": round(nav, 8),
            })
    if nav_match:
        disclosed_nav = float(nav_match.group(1))
        if disclosed_nav <= 0 or abs(nav - disclosed_nav) / disclosed_nav > 0.025:
            return []
    return _valid_nav_points(points)


def _looks_like_pdf_product_curve(name: str, file_stem: str) -> bool:
    """Reject benchmark, sector and chart-label curves before creating products."""
    token = name.replace(" ", "").strip()
    if len(token) < 3:
        return False
    generic = (
        "产品净值", "累计净值", "单位净值", "图例", "超额收益", "基准", "指数",
        "沪深300", "中证", "上证", "国债", "南华", "商品指数", "板块", "行业",
        "农产品", "黑色", "有色", "能化", "贵金属", "策略收益", "组合收益",
    )
    if any(label in token for label in generic):
        return False
    # A model returning the document title instead of a legend has not
    # established a concrete fund identity.  Accept it only when the filename
    # itself looks like a product name under the same conservative gate.
    if token == file_stem.replace(" ", "").strip():
        return _looks_like_product_name(file_stem)
    return _looks_like_product_name(name)


def _auto_curve_colors_from_pdf(pdf_bytes: bytes) -> list[str]:
    """Render first PDF page and auto-detect curve color for CV tracing."""
    from app.services.chart_extractor.pdf_renderer import render_pdf_pages

    try:
        render_result = render_pdf_pages(pdf_bytes, max_pages=1)
    except Exception:
        return []
    if not render_result.pages:
        return []
    return _auto_curve_colors(render_result.pages[0].image_bytes)


def _find_or_create_product(
    session: Any,
    name: str,
    strategy: str | None,
    inception: date | None,
    close: date | None,
    manager_name: str | None = None,
) -> ProductEntity:
    cleaned_name = name.strip() or "未命名产品"
    product = session.execute(select(ProductEntity).where(ProductEntity.standard_name == cleaned_name)).scalars().first()
    if product is None:
        product = store.create_product(
            session,
            standard_name=cleaned_name,
            manager_name=manager_name,
            strategy=strategy,
            inception_date=inception,
        )
    elif strategy and not product.strategy:
        product.strategy = strategy
    if manager_name and not product.manager_name:
        product.manager_name = manager_name
    if close and (product.close_date is None or close > product.close_date):
        product.close_date = close
    session.commit()
    session.refresh(product)
    return product


def _find_column(frame: pd.DataFrame, candidates: tuple[str, ...]) -> Any | None:
    normalized = {str(column).strip().lower().replace(" ", ""): column for column in frame.columns}
    for candidate in candidates:
        found = normalized.get(candidate.lower().replace(" ", ""))
        if found is not None:
            return found
    return None


def _infer_frequency(dates: list[date], *, pixel_trace: bool = False) -> str:
    if len(dates) < 2:
        return "unknown"
    # Pixel tracing emits one point per horizontal pixel.  Those points are
    # interpolation samples, not daily observations; treating them as daily
    # would inflate volatility and Sharpe.  Unless VLM/user metadata provides
    # another frequency, use the product-report default of weekly.
    if pixel_trace and len(dates) >= 30:
        return "weekly"
    gaps = [(right - left).days for left, right in zip(dates, dates[1:])]
    median_gap = sorted(gaps)[len(gaps) // 2]
    if median_gap <= 3:
        return "daily"
    if median_gap <= 10:
        return "weekly"
    return "monthly"
