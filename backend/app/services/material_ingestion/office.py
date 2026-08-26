"""Office document ingestion (DOCX/PPTX text and disclosed tables)."""

from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd

from app.services import product_store as store

from .helpers import _find_or_create_product
from .table import _ingest_nav_frame


def _ingest_office_document(session: Any, file_id: str, content: bytes, filename: str, suffix: str) -> int:
    """Extract Word/PPT text and disclosed tables without inventing chart data."""
    product = _find_or_create_product(session, Path(filename).stem, None, None, None)
    parts = _extract_office_parts(content, suffix)
    if not any(part["text"] or part["tables"] for part in parts):
        raise ValueError("文档未提取到可读取的文本或表格")

    total_rows = 0
    for part in parts:
        location = part["location"]
        if part["text"]:
            store.add_fragment(
                session,
                file_id=file_id,
                fragment_type="text",
                page_number=part["page_number"],
                content_text=part["text"][:20000],
                content_data={"location": location, "method": part["method"], "review_status": "pending"},
                ocr_confidence=1.0,
                product_id=product.id,
            )
        for table_index, table in enumerate(part["tables"], 1):
            if len(table) < 2:
                continue
            header, *rows = table
            if not header or not rows:
                continue
            header_tokens = [str(value).replace(" ", "").lower() for value in header]
            has_explicit_date = any(any(label in token for label in ("date", "日期", "净值日期", "估值日期")) for token in header_tokens)
            has_explicit_nav = any(any(label in token for label in ("nav", "净值", "单位净值", "累计净值", "acc_nav")) for token in header_tokens)
            if not (has_explicit_date and has_explicit_nav):
                continue
            frame = pd.DataFrame(rows, columns=header)
            total_rows += _ingest_nav_frame(
                session,
                file_id,
                frame,
                filename,
                location=f"{location}; table={table_index}",
            )
    return total_rows


def _extract_office_parts(content: bytes, suffix: str) -> list[dict[str, Any]]:
    """Return text/table evidence with stable document locations for Office files."""
    if suffix == ".docx":
        from docx import Document

        document = Document(BytesIO(content))
        text = "\n".join(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
        tables = [[ [cell.text.strip() for cell in row.cells] for row in table.rows ] for table in document.tables]
        return [{"page_number": None, "location": "Word document", "text": text, "tables": tables, "method": "DOCX XML 提取"}]

    parts: list[dict[str, Any]] = []
    # Read slide XML directly. python-pptx materializes embedded media and can
    # consume gigabytes on image-heavy manager decks even though ingestion only
    # needs text/table evidence.
    from lxml import etree

    with zipfile.ZipFile(BytesIO(content)) as archive:
        slide_names = sorted(
            (
                name for name in archive.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            ),
            key=lambda name: int(re.search(r"(\d+)\.xml$", name).group(1)),
        )
        for index, slide_name in enumerate(slide_names, 1):
            root = etree.fromstring(archive.read(slide_name))
            namespaces = {
                "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
            }
            text_values = [value.strip() for value in root.xpath("//a:t/text()", namespaces=namespaces) if value.strip()]
            tables: list[list[list[str]]] = []
            for table in root.xpath("//a:tbl", namespaces=namespaces):
                rows: list[list[str]] = []
                for row in table.xpath("./a:tr", namespaces=namespaces):
                    cells = ["".join(cell.xpath(".//a:t/text()", namespaces=namespaces)).strip() for cell in row.xpath("./a:tc", namespaces=namespaces)]
                    rows.append(cells)
                if rows:
                    tables.append(rows)
            parts.append({"page_number": index, "location": f"slide={index}", "text": "\n".join(text_values), "tables": tables, "method": "PPTX XML 流式提取"})
    return parts


