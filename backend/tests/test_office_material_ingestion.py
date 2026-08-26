from io import BytesIO

from docx import Document
from pptx import Presentation

from app.services.material_ingestion import _extract_office_parts


def test_extracts_word_text_and_table() -> None:
    document = Document()
    document.add_paragraph("CTA 产品说明")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "日期", "净值"
    table.cell(1, 0).text, table.cell(1, 1).text = "2026-01-02", "1.01"
    table.cell(2, 0).text, table.cell(2, 1).text = "2026-01-09", "1.03"
    output = BytesIO()
    document.save(output)

    parts = _extract_office_parts(output.getvalue(), ".docx")

    assert parts[0]["location"] == "Word document"
    assert "CTA 产品说明" in parts[0]["text"]
    assert parts[0]["tables"][0][0] == ["日期", "净值"]


def test_extracts_ppt_slide_text_and_table() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.add_textbox(0, 0, 3000000, 500000).text_frame.text = "商品 CTA 策略"
    table = slide.shapes.add_table(3, 2, 0, 600000, 3000000, 1000000).table
    table.cell(0, 0).text, table.cell(0, 1).text = "日期", "净值"
    table.cell(1, 0).text, table.cell(1, 1).text = "2026-01-02", "1.01"
    table.cell(2, 0).text, table.cell(2, 1).text = "2026-01-09", "1.03"
    output = BytesIO()
    presentation.save(output)

    parts = _extract_office_parts(output.getvalue(), ".pptx")

    assert parts[0]["page_number"] == 1
    assert parts[0]["location"] == "slide=1"
    assert "商品 CTA 策略" in parts[0]["text"]
    assert parts[0]["tables"][0][0] == ["日期", "净值"]
